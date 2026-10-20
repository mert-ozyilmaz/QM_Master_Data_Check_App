"""Excel dosyalarını okuyup kuralların kullanacağı nesnelere çevirir."""
from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import load_workbook

from . import ayarlar
from .yardimcilar import dolu, kolon_harfi, kolon_no, metin, norm_ad, sayi, x_mi


class Satir:
    """Bir veri satırı. satir["B"] -> ham değer; satir.n -> Excel satır numarası."""
    __slots__ = ("n", "_v", "_ws")

    def __init__(self, n, degerler, ws):
        self.n, self._v, self._ws = n, degerler, ws

    def __getitem__(self, kolon: str):
        i = kolon_no(kolon)
        return self._v[i] if i < len(self._v) else None

    def metin(self, kolon: str) -> str:
        return metin(self[kolon])

    def dolu(self, kolon: str) -> bool:
        return dolu(self[kolon])

    def sayi(self, kolon: str):
        return sayi(self[kolon])

    def x(self, kolon: str) -> bool:
        return x_mi(self[kolon])

    def bicim(self, kolon: str) -> str:
        """Hücrenin sayı biçimi (ör. '0.00', 'General')."""
        return self._ws[f"{kolon}{self.n}"].number_format

    def hucreler(self):
        """Dolu hücreleri (kolon_harfi, değer) olarak verir."""
        for i, v in enumerate(self._v):
            if dolu(v):
                yield kolon_harfi(i), v


@dataclass
class Sayfa:
    ad: str
    dosya: str                      # standart dosya adı (ör. 'Ana kontrol karakteristikleri')
    baslik_satiri: int
    bilgi: bool = False             # True ise yalnızca bilgi sayfası: kurallar çalışmaz
    satirlar: list = field(default_factory=list)   # yalnızca dolu veri satırları

    def kolon_degerleri(self, kolon: str) -> set:
        return {s.metin(kolon) for s in self.satirlar if s.dolu(kolon)}


@dataclass
class Dosya:
    ad: str                         # standart ad
    yuklenen_ad: str
    sayfalar: list = field(default_factory=list)

    def sayfa(self, ad: str):
        n = norm_ad(ad)
        return next((s for s in self.sayfalar if norm_ad(s.ad) == n), None)


class Calisma:
    """Yüklenen tüm dosyalar."""
    def __init__(self):
        self.dosyalar: dict[str, Dosya] = {}

    def dosya(self, ad: str):
        return self.dosyalar.get(norm_ad(ad))


def dosya_yukle(icerik: bytes, yuklenen_ad: str):
    """(Dosya | None, hata_mesaji) döner."""
    standart = ayarlar.standart_dosya_adi(yuklenen_ad)
    if standart is None:
        return None, "Tanınmayan dosya adı"
    try:
        wb = load_workbook(BytesIO(icerik), data_only=True)
    except Exception as e:  # bozuk dosya
        return None, f"Dosya açılamadı: {e}"
    dosya = Dosya(ad=standart, yuklenen_ad=yuklenen_ad)
    for ws in wb.worksheets:
        bilgi = ayarlar.bilgi_sayfasi_mi(standart, ws.title)
        bas = ayarlar.baslik_satiri_sayisi(standart, ws.title)
        sayfa = Sayfa(ad=ws.title.strip(), dosya=standart, baslik_satiri=bas, bilgi=bilgi)
        if not bilgi:
            for n, vals in enumerate(ws.iter_rows(values_only=True), start=1):
                if n <= bas:
                    continue                       # Madde 1.1: başlıklar kontrol edilmez
                if any(dolu(v) for v in vals):     # tamamen boş satırlar atlanır
                    sayfa.satirlar.append(Satir(n, vals, ws))
        dosya.sayfalar.append(sayfa)
    return dosya, ""
