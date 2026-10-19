"""Word (.docx) rapor üretimi."""
from collections import Counter, OrderedDict
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .yardimcilar import madde_sirasi


def _golge(hucre, renk):
    tc_pr = hucre._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), renk)
    tc_pr.append(shd)


def _baslik_tekrarla(satir):
    tr_pr = satir._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _tablo(doc, basliklar, satirlar, genislikler_cm):
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, b in enumerate(basliklar):
        h = t.rows[0].cells[i]
        h.text = ""
        r = h.paragraphs[0].add_run(b)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(255, 255, 255)
        _golge(h, "1F3864")
    _baslik_tekrarla(t.rows[0])
    for satir in satirlar:
        cells = t.add_row().cells
        for i, v in enumerate(satir):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(8.5)
    for row in t.rows:
        for i, w in enumerate(genislikler_cm):
            row.cells[i].width = Cm(w)
    # Word/LibreOffice'in genişlikleri uygulaması için tablo ızgarasını da ayarla
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), genislikler_cm):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    return t


def rapor_olustur(sonuc) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(1.5))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading("QM Ana Veri Kontrol Raporu", 0)
    doc.add_paragraph(f"Rapor tarihi: {sonuc.tarih:%d.%m.%Y %H:%M}")
    doc.add_paragraph("Kontroller, 'QM Master Data Rules' dokümanındaki kurallara göre yapılmıştır. "
                      "Satır numaraları Excel'deki satır numaralarıdır.")

    # ---------------- 1. ÖZET ----------------
    doc.add_heading("1. Özet", 1)
    toplam = len(sonuc.bulgular)
    doc.add_paragraph(f"Kontrol edilen dosya sayısı: {len(sonuc.yuklenen_dosyalar)}    |    "
                      f"Çalıştırılan kural sayısı: {len(sonuc.calisan_kurallar)}    |    "
                      f"Toplam bulgu: {toplam}")
    if toplam == 0:
        p = doc.add_paragraph()
        r = p.add_run("Kontrol edilen kurallara göre hata bulunamadı.")
        r.bold = True

    doc.add_heading("1.1 Dosya / sayfa bazında özet", 2)
    sayac = Counter((b.dosya, b.sayfa) for b in sonuc.bulgular)
    _tablo(doc, ["Dosya", "Sayfa", "Kontrol edilen satır", "Bulgu sayısı"],
           [(d, s, n, sayac.get((d, s), 0)) for d, s, n in sonuc.kontrol_edilen],
           [8, 8, 5, 4])

    doc.add_heading("1.2 Kural bazında özet", 2)
    kural_sayac = Counter(b.kural_id for b in sonuc.bulgular)
    _tablo(doc, ["Madde", "Kapsam", "Kural", "Bulgu sayısı"],
           [(k.id, k.kapsam, k.baslik, kural_sayac.get(k.id, 0)) for k in sonuc.calisan_kurallar],
           [2, 2, 19, 3.5])

    # ---------------- 2. DETAY ----------------
    doc.add_heading("2. Hata detayları", 1)
    if toplam == 0:
        doc.add_paragraph("Hata bulunamadı.")
    gruplar = OrderedDict()
    for b in sonuc.bulgular:
        gruplar.setdefault(b.dosya, OrderedDict()).setdefault(b.sayfa, []).append(b)
    for dosya, sayfalar in gruplar.items():
        doc.add_heading(dosya, 2)
        for sayfa, liste in sayfalar.items():
            doc.add_heading(f"Sayfa: {sayfa}  ({len(liste)} bulgu)", 3)
            _tablo(doc, ["Satır", "Hücre", "Değer", "Madde", "Açıklama"],
                   [(b.satir, b.kolon, (b.deger[:60] + "…") if len(b.deger) > 60 else b.deger,
                     b.kural_id, b.mesaj) for b in liste],
                   [1.5, 2.2, 4.5, 1.8, 16.7])

    # ---------------- 3. EKLER ----------------
    doc.add_heading("3. Ekler ve notlar", 1)
    no = iter(range(1, 10))

    if sonuc.atlananlar:
        doc.add_heading(f"3.{next(no)} Kontrol edilemeyen kurallar", 2)
        doc.add_paragraph("Aşağıdaki kurallar, gerekli referans dosya/sayfa bulunamadığı veya çalışırken "
                          "hata oluştuğu için (kısmen) uygulanamadı:")
        _tablo(doc, ["Madde", "Neden"], sonuc.atlananlar, [2, 24.7])

    if sonuc.taninmayan_dosyalar or sonuc.yuklenmeyen_dosyalar:
        doc.add_heading(f"3.{next(no)} Dosya durumu", 2)
        if sonuc.taninmayan_dosyalar:
            doc.add_paragraph("Tanınmayan / okunamayan dosyalar (kontrol edilmedi):")
            _tablo(doc, ["Dosya", "Neden"], sonuc.taninmayan_dosyalar, [12, 14.7])
        if sonuc.yuklenmeyen_dosyalar:
            doc.add_paragraph("Yüklenmeyen dosyalar: " + ", ".join(sonuc.yuklenmeyen_dosyalar))

    if sonuc.pasif_kurallar:
        doc.add_heading(f"3.{next(no)} Pasif kurallar", 2)
        doc.add_paragraph("Aşağıdaki kurallar kodda tanımlı ancak pasif olduğu için çalıştırılmadı:")
        _tablo(doc, ["Madde", "Kural", "Açıklama"],
               [(k.id, k.baslik, k.varsayim) for k in sonuc.pasif_kurallar], [2, 10, 14.7])

    varsayimli = [k for k in sonuc.calisan_kurallar if k.varsayim]
    if varsayimli:
        doc.add_heading(f"3.{next(no)} Yorumlanan / varsayım yapılan kurallar", 2)
        _tablo(doc, ["Madde", "Varsayım / not"],
               [(k.id, k.varsayim) for k in sorted(varsayimli, key=lambda k: madde_sirasi(k.id))], [2, 24.7])

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
